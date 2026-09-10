from io import BytesIO
from types import SimpleNamespace

import numpy as np
import pytest
import xlsxwriter
from openpyxl import load_workbook

from manic.processors.integration import (
    calculate_peak_areas,
    compute_baseline_area,
)
from manic.processors.natural_abundance_correction import NaturalAbundanceCorrector
from manic.sheet_generators import abundances


def test_c1_round_trip_recovers_true_vector():
    corrector = NaturalAbundanceCorrector()
    matrix = corrector.build_correction_matrix("C1", "C", 1)
    true = np.array([[80.0], [20.0]])
    measured = matrix @ true
    got = corrector._correct_vectorized_direct(measured, matrix)
    np.testing.assert_allclose(
        got,
        true,
        atol=1e-9,
        err_msg="solve(C, C @ [[80],[20]]) == [[80],[20]] for C1 label_atoms=1",
    )


def test_tbdms_c3h6o3_round_trip_recovers_true_vector():
    corrector = NaturalAbundanceCorrector()
    deriv_formula, _ = corrector.calculate_derivative_formula("C3H6O3", tbdms=1)
    matrix = corrector.build_correction_matrix(deriv_formula, "C", 3)
    true = np.array([[50.0], [10.0], [5.0], [35.0]])
    measured = matrix @ true
    got = corrector._correct_vectorized_direct(measured, matrix)
    np.testing.assert_allclose(
        got,
        true,
        atol=1e-9,
        err_msg="solve(C, C @ [[50],[10],[5],[35]]) == true for C3H6O3 tbdms=1 label_atoms=3",
    )


def test_unlabelled_1x1_solve_divides_by_diagonal():
    matrix = np.array([[0.9893]])
    measured = np.array([[100.0]])
    got = NaturalAbundanceCorrector()._correct_vectorized_direct(measured, matrix)
    expected = 100.0 / 0.9893
    assert got[0, 0] == pytest.approx(expected, abs=1e-12), (
        "1x1 solve([[0.9893]], [[100]]) = 100/0.9893 = 101.08157262710957"
    )


def test_baseline_area_clamps_negative_line_at_zero():
    time = np.array([0.0, 1.0, 2.0, 3.0, 4.0, 5.0, 6.0, 7.0, 8.0, 9.0])
    intensity = np.array([0.0, 0.0, 0.0, 100.0, 100.0, 100.0, 100.0, 30.0, 20.0, 10.0])
    area = compute_baseline_area(time, intensity)
    old_area = 90.0
    new_area = 90.21731748726658
    assert area == pytest.approx(new_area, abs=1e-9), (
        "polyfit first3=[0,0,0] last3=[30,20,10] on t=[0..9] gives "
        "B(0)=-1.0322580645161248, B(9)=21.032258064516135, width=9. "
        "old 0.5*(-1.0322580645161248+21.032258064516135)*9 = 90.0. "
        "new 0.5*21.032258064516135**2 / 22.06451612903226 * 9 = 90.21731748726658 "
        "(integral of max(line, 0))"
    )
    assert area != pytest.approx(old_area, abs=1e-6)

    uncorrected = calculate_peak_areas(
        time, intensity, label_atoms=0, baseline_correction=False
    )[0]
    corrected = calculate_peak_areas(
        time, intensity, label_atoms=0, baseline_correction=True
    )[0]
    assert corrected <= uncorrected, (
        "baseline can only reduce area: "
        f"uncorrected={uncorrected}, corrected={corrected}"
    )


def test_parse_formula_keeps_two_letter_chlorine():
    parsed = NaturalAbundanceCorrector().parse_formula("C2H3Cl")
    assert parsed == {"C": 2, "H": 3, "Cl": 1}, (
        "C2H3Cl parses as C=2 H=3 Cl=1, not an unknown element l"
    )
    assert "l" not in parsed


def test_c1cl1_column0_is_carbon_convolved_with_chlorine():
    matrix = NaturalAbundanceCorrector().build_correction_matrix(
        "C1Cl1", "C", 1, max_isotopologues=3
    )
    expected = np.array([0.74949, 0.00811, 0.23981])
    np.testing.assert_allclose(
        matrix[:3, 0],
        expected,
        atol=5e-6,
        err_msg=(
            "C1Cl1 column 0 is C[0.9893, 0.0107] conv Cl[0.7576, 0, 0.2424] "
            "truncated to 3: [0.9893*0.7576, 0.0107*0.7576, 0.9893*0.2424] "
            "= [0.74949, 0.00811, 0.23981]. label_atoms=1 because C1 has one "
            "carbon; max_isotopologues=3 keeps the Cl +2 channel"
        ),
    )


def test_assumed_mrrf_writes_relative_unit():
    class Provider:
        def get_all_compounds(self):
            return [
                {
                    "compound_name": "Calibrated",
                    "mass0": 100.0,
                    "retention_time": 1.0,
                    "amount_in_std_mix": 1.0,
                    "mm_files": "*MM*",
                    "baseline_correction": 1,
                },
                {
                    "compound_name": "Assumed",
                    "mass0": 110.0,
                    "retention_time": 2.0,
                    "amount_in_std_mix": 1.0,
                    "mm_files": "*NONE*",
                    "baseline_correction": 1,
                },
                {
                    "compound_name": "ISTD",
                    "mass0": 90.0,
                    "retention_time": 0.5,
                    "amount_in_std_mix": 1.0,
                    "int_std_amount": 1.0,
                    "mm_files": "*MM*",
                    "baseline_correction": 1,
                },
            ]

        def get_all_samples(self):
            return ["S1"]

        def resolve_mm_samples(self, mm_field):
            if mm_field == "*MM*":
                return ["MM1"]
            return []

        def get_sample_corrected_data(self, sample_name):
            return {
                "Calibrated": [100.0],
                "Assumed": [80.0],
                "ISTD": [50.0],
            }

        def get_mrrf_values(
            self,
            compounds,
            internal_standard_compound,
            internal_standard_isotope_index=0,
            assumed=None,
        ):
            values = {}
            for row in compounds:
                name = row["compound_name"]
                if name == internal_standard_compound:
                    values[name] = 1.0
                    continue
                if self.resolve_mm_samples(row.get("mm_files")):
                    values[name] = 2.0
                else:
                    values[name] = 1.0
                    if assumed is not None:
                        assumed.add(name)
            return values

    output = BytesIO()
    workbook = xlsxwriter.Workbook(output, {"in_memory": True})
    abundances.write(
        workbook,
        SimpleNamespace(
            internal_standard_compound="ISTD",
            internal_standard_reference_isotope=0,
        ),
        None,
        0,
        100,
        provider=Provider(),
    )
    workbook.close()
    output.seek(0)
    sheet = load_workbook(output)["Abundances"]
    assert sheet["C5"].value == "nmol", (
        "Calibrated has amount_in_std_mix=1.0 and MM samples, so Units is nmol"
    )
    assert sheet["D5"].value == "Relative", (
        "Assumed has amount_in_std_mix=1.0 but resolve_mm_samples returns none, "
        "so MRRF is assumed and Units is Relative"
    )
