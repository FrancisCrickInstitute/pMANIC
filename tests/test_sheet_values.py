import zlib
from io import BytesIO
from types import SimpleNamespace

import numpy as np
import pytest
import xlsxwriter
from openpyxl import load_workbook

from manic.io.data_exporter import DataExporter
from manic.io.data_provider import DataProvider
from manic.io.in_memory_provider import InMemoryDataProvider
from manic.models import database
from manic.models.analysis import AnalysisMode
from manic.processors.calibration import calculate_mrrf_values
from manic.processors.eic_correction_manager import process_all_corrections
from manic.processors.natural_abundance_correction import NaturalAbundanceCorrector
from manic.sheet_generators import (
    abundances,
    corrected_values,
    isotope_ratios,
    label_incorporation,
    unlabelled_results,
)


def _sheet_cell(write_fn, provider, sheet_name, cell):
    output = BytesIO()
    workbook = xlsxwriter.Workbook(output, {"in_memory": True})
    write_fn(workbook, SimpleNamespace(), None, 0, 100, provider=provider)
    workbook.close()
    output.seek(0)
    return load_workbook(output)[sheet_name][cell].value


def test_label_incorporation_background_subtracted_percentage():
    class Provider:
        def get_all_compounds(self):
            return [
                {
                    "compound_name": "Pyruvate",
                    "mass0": 174.0,
                    "retention_time": 6.0,
                    "label_atoms": 1,
                    "baseline_correction": 1,
                    "mm_files": "*MM*",
                }
            ]

        def get_all_samples(self):
            return ["S1"]

        def get_background_ratios(self, compounds):
            return {"Pyruvate": 0.05}

        def get_sample_corrected_data(self, sample_name):
            return {"Pyruvate": [100.0, 25.0]}

    value = _sheet_cell(
        label_incorporation.write, Provider(), "% Label Incorporation", "C5"
    )
    assert value == pytest.approx(16.0), (
        "(25 - 0.05*100) / 125 * 100"
    )


def test_label_incorporation_zero_total_is_zero():
    class Provider:
        def get_all_compounds(self):
            return [
                {
                    "compound_name": "Pyruvate",
                    "mass0": 174.0,
                    "retention_time": 6.0,
                    "label_atoms": 1,
                    "baseline_correction": 1,
                    "mm_files": None,
                }
            ]

        def get_all_samples(self):
            return ["S1"]

        def get_background_ratios(self, compounds):
            return {"Pyruvate": 0.05}

        def get_sample_corrected_data(self, sample_name):
            return {"Pyruvate": [0.0, 0.0]}

    value = _sheet_cell(
        label_incorporation.write, Provider(), "% Label Incorporation", "C5"
    )
    assert value == 0.0, "total_signal==0 yields 0.0"


def test_isotope_ratio_uses_sum_of_areas_as_denominator():
    class Provider:
        def get_all_compounds(self):
            return [
                {
                    "compound_name": "Pyruvate",
                    "label_atoms": 2,
                    "mass0": 174.0,
                    "retention_time": 6.0,
                    "baseline_correction": 1,
                }
            ]

        def get_all_samples(self):
            return ["S1"]

        def get_sample_corrected_data(self, sample_name):
            return {"Pyruvate": [100.0, 50.0, 25.0]}

    m0 = _sheet_cell(isotope_ratios.write, Provider(), "Isotope Ratio", "C5")
    m1 = _sheet_cell(isotope_ratios.write, Provider(), "Isotope Ratio", "D5")
    m2 = _sheet_cell(isotope_ratios.write, Provider(), "Isotope Ratio", "E5")
    assert m0 == pytest.approx(100.0 / 175.0), "100 / (100+50+25)"
    assert m1 == pytest.approx(50.0 / 175.0), "50 / (100+50+25)"
    assert m2 == pytest.approx(25.0 / 175.0), "25 / (100+50+25)"


def test_isotope_ratio_all_zero_is_zero():
    class Provider:
        def get_all_compounds(self):
            return [
                {
                    "compound_name": "Pyruvate",
                    "label_atoms": 2,
                    "mass0": 174.0,
                    "retention_time": 6.0,
                    "baseline_correction": 1,
                }
            ]

        def get_all_samples(self):
            return ["S1"]

        def get_sample_corrected_data(self, sample_name):
            return {"Pyruvate": [0.0, 0.0, 0.0]}

    m0 = _sheet_cell(isotope_ratios.write, Provider(), "Isotope Ratio", "C5")
    assert m0 == 0.0, "total_area==0 writes 0.0"


def test_corrected_values_sheet_writes_stored_vector_unchanged():
    stored = [12.5, 3.25]

    class Provider:
        def get_all_compounds(self):
            return [
                {
                    "compound_name": "Pyruvate",
                    "label_atoms": 1,
                    "mass0": 174.0,
                    "retention_time": 6.0,
                    "baseline_correction": 1,
                }
            ]

        def get_all_samples(self):
            return ["S1"]

        def get_sample_corrected_data(self, sample_name):
            return {"Pyruvate": list(stored)}

    m0 = _sheet_cell(corrected_values.write, Provider(), "Corrected Values", "C5")
    m1 = _sheet_cell(corrected_values.write, Provider(), "Corrected Values", "D5")
    assert m0 == 12.5
    assert m1 == 3.25


def test_build_correction_matrix_tbdms_c3h6o3_m0_to_m0():
    corrector = NaturalAbundanceCorrector()
    formula, counts = corrector.calculate_derivative_formula("C3H6O3", tbdms=1)
    assert counts["C"] == 5
    assert counts["H"] == 11
    assert counts["O"] == 3
    assert counts["Si"] == 1
    matrix = corrector.build_correction_matrix(formula, "C", 3)
    assert round(float(matrix[0, 0]), 3) == 0.866


def test_calculate_mrrf_two_mm_files_is_2_5(empty_db):
    with database.get_connection() as conn:
        conn.execute(
            "INSERT INTO compounds (compound_name, retention_time, mass0, label_atoms, "
            "amount_in_std_mix, mm_files) VALUES (?, 1.0, 100.0, 0, 1.0, '*MM*')",
            ("Target",),
        )
        conn.execute(
            "INSERT INTO compounds (compound_name, retention_time, mass0, label_atoms, "
            "amount_in_std_mix, mm_files) VALUES (?, 2.0, 200.0, 0, 1.0, '*MM*')",
            ("ISTD",),
        )

    class Provider:
        def resolve_mm_samples(self, pattern):
            return ["MM1", "MM2"]

        def get_sample_corrected_data(self, sample_name):
            return {"Target": [250.0], "ISTD": [100.0]}

        def get_compound_total_area(self, sample_name, compound_name):
            return 250.0 if compound_name == "Target" else 100.0

    compounds = [
        {"compound_name": "Target", "amount_in_std_mix": 1.0, "mm_files": "*MM*"},
        {"compound_name": "ISTD", "amount_in_std_mix": 1.0, "mm_files": "*MM*"},
    ]
    values = calculate_mrrf_values(Provider(), compounds, "ISTD")
    assert values["Target"] == pytest.approx(2.5), "(250/1) / (100/1)"


def _zero_signal_compounds():
    return [
        {"compound_name": "Silent", "amount_in_std_mix": 1.0, "mm_files": "*MM*"},
        {"compound_name": "ISTD", "amount_in_std_mix": 1.0, "mm_files": "*MM*"},
    ]


def test_calculate_mrrf_zero_metabolite_signal_is_assumed_one():
    class Provider:
        def resolve_mm_samples(self, pattern):
            return ["MM1"]

        def get_sample_corrected_data(self, sample_name):
            return {"Silent": [0.0], "ISTD": [100.0]}

        def get_compound_total_area(self, sample_name, compound_name):
            return 0.0 if compound_name == "Silent" else 100.0

    assumed = set()
    values = calculate_mrrf_values(Provider(), _zero_signal_compounds(), "ISTD", assumed=assumed)
    assert values["Silent"] == 1.0
    assert assumed == {"Silent"}


def test_in_memory_mrrf_zero_metabolite_signal_is_assumed_one():
    provider = InMemoryDataProvider(
        _zero_signal_compounds(), ["MM1"], {"MM1": {"Silent": [0.0], "ISTD": [100.0]}}
    )
    assumed = set()
    values = provider.get_mrrf_values(_zero_signal_compounds(), "ISTD", assumed=assumed)
    assert values["Silent"] == 1.0
    assert assumed == {"Silent"}


def test_unlabelled_abundances_units_row_marks_assumed_mrrf_as_relative():
    class Provider:
        def get_mrrf_values(self, compounds, internal_standard, internal_standard_isotope_index=0, assumed=None):
            assumed.add("NoStandard")
            return {"ISTD": 1.0, "NoStandard": 1.0}

        def resolve_mm_samples(self, pattern):
            return []

    exporter = SimpleNamespace(_provider=Provider(), internal_standard_compound="ISTD")
    compounds = [
        SimpleNamespace(
            compound_name=name,
            baseline_correction=True,
            retention_time=1.0,
            analysis_channels=[SimpleNamespace(mz=100.0)],
        )
        for name in ("ISTD", "NoStandard")
    ]
    meta = {
        name: {"amount_in_std_mix": 1.0, "int_std_amount": 1.0, "mm_files": "*MM*"}
        for name in ("ISTD", "NoStandard")
    }
    output = BytesIO()
    workbook = xlsxwriter.Workbook(output, {"in_memory": True})
    unlabelled_results._write_abundances(
        workbook, exporter, [], compounds, list(meta.values()), meta, {}, None
    )
    workbook.close()
    output.seek(0)
    sheet = load_workbook(output)["Abundances"]
    assert [sheet.cell(4, col).value for col in (3, 4)] == ["nmol", "Relative"]
    assert exporter.assumed_mrrf == {"NoStandard"}


def test_process_all_corrections_matches_correct_time_series(empty_db):
    time = np.array([0.0, 1.0, 2.0], dtype=np.float64)
    intensity = np.array(
        [
            [100.0, 80.0, 20.0],
            [10.0, 8.0, 2.0],
        ],
        dtype=np.float64,
    )
    with database.get_connection() as conn:
        conn.execute(
            "INSERT INTO compounds (compound_name, retention_time, mass0, label_atoms, "
            "formula, label_type, tbdms, meox, me) "
            "VALUES ('Pyruvate', 1.0, 87.0, 1, 'C1H4', 'C', 0, 0, 0)"
        )
        conn.execute("INSERT INTO samples (sample_name, file_name) VALUES ('S1', '/fake/S1.cdf')")
        conn.execute(
            "INSERT INTO eic (sample_name, compound_name, x_axis, y_axis, rt_window) "
            "VALUES (?, ?, ?, ?, 0.4)",
            (
                "S1",
                "Pyruvate",
                zlib.compress(time.tobytes()),
                zlib.compress(intensity.ravel().tobytes()),
            ),
        )

    assert process_all_corrections() == 1
    expected = NaturalAbundanceCorrector().correct_time_series(
        intensity, "C1H4", "C", 1, 0, 0, 0
    )
    with database.get_connection() as conn:
        row = conn.execute(
            "SELECT y_axis_corrected FROM eic_corrected "
            "WHERE sample_name = 'S1' AND compound_name = 'Pyruvate'"
        ).fetchone()
    stored = np.frombuffer(zlib.decompress(row["y_axis_corrected"]), dtype=np.float64)
    np.testing.assert_allclose(stored, expected.ravel())


def test_export_to_excel_stops_and_leaves_no_file_when_callback_says_stop(
    empty_db, tmp_path, monkeypatch
):
    monkeypatch.setattr(
        DataProvider,
        "load_bulk_sample_data",
        lambda self, progress_callback=None: progress_callback(10) or {},
    )
    export_path = tmp_path / "cancelled.xlsx"
    exporter = DataExporter(AnalysisMode.LABELLED)
    assert exporter.export_to_excel(str(export_path), lambda _value: False) is False
    assert not export_path.exists()


def test_abundances_units_row_marks_assumed_mrrf_as_relative(empty_db):
    with database.get_connection() as conn:
        conn.execute(
            "INSERT INTO compounds (compound_name, formula, mass0, retention_time, "
            "label_atoms, amount_in_std_mix, int_std_amount, mm_files) VALUES "
            "('ISTD', 'C1', 100.0, 5.0, 1, 1.0, 1.0, '*MM*'), "
            "('NoStandard', 'C1', 300.0, 7.0, 1, 1.0, 1.0, '*Nothing*')"
        )
    exporter = DataExporter(AnalysisMode.LABELLED)
    exporter.set_internal_standard("ISTD")
    output = BytesIO()
    workbook = xlsxwriter.Workbook(output, {"in_memory": True})
    abundances.write(workbook, exporter, None, 0, 100)
    workbook.close()
    output.seek(0)
    sheet = load_workbook(output)["Abundances"]
    units = {sheet.cell(1, col).value: sheet.cell(5, col).value for col in range(3, 5)}
    assert units == {"ISTD": "nmol", "NoStandard": "Relative"}
    assert exporter.assumed_mrrf == {"NoStandard"}
