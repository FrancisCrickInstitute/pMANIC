from __future__ import annotations

from openpyxl import load_workbook

from manic.io.data_exporter import DataExporter


def test_export(case, db, benchmark, repeat, tmp_path):
    def export():
        exporter = DataExporter(case.mode)
        exporter.set_internal_standard(case.internal_standard)
        exporter.set_use_legacy_integration(False)
        path = tmp_path / "out.xlsx"
        exporter.export_to_excel(path, include_carbon_enrichment=case.labelled)
        return path

    path = benchmark.pedantic(export, rounds=repeat)
    wb = load_workbook(path, read_only=True)
    assert {"Raw Values", "Abundances"} <= set(wb.sheetnames)
    rows = list(wb["Abundances"].iter_rows(min_col=2, max_col=2, values_only=True))
    samples = {r[0] for r in rows if r[0] and str(r[0]).startswith(("Bio_", "Std_", "MM_"))}
    assert len(samples) == case.n_samples
    wb.close()
