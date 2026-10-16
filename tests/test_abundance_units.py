from unittest.mock import MagicMock

import pytest

from manic.sheet_generators.abundances import write


class TestAbundanceExportLogic:
    @pytest.fixture
    def setup_base_mocks(self):
        workbook = MagicMock()
        worksheet = MagicMock()
        workbook.add_worksheet.return_value = worksheet

        provider = MagicMock()
        provider.get_all_samples.return_value = ["Sample1"]
        provider.get_sample_corrected_data.return_value = {
            "ISTD": [1000.0],
            "Cmp_Absolute": [500.0],
            "Cmp_Relative": [500.0],
        }
        provider.get_all_compounds.return_value = [
            {
                "compound_name": "ISTD",
                "amount_in_std_mix": 1.0,
                "int_std_amount": 10.0,
                "mass0": 100,
                "retention_time": 5.0,
            },
            {
                "compound_name": "Cmp_Absolute",
                "amount_in_std_mix": 5.0,
                "mass0": 110,
                "retention_time": 5.5,
            },
            {
                "compound_name": "Cmp_Relative",
                "amount_in_std_mix": 0.0,
                "mass0": 120,
                "retention_time": 6.0,
            },
        ]
        provider.get_mrrf_values.return_value = {"Cmp_Absolute": 2.0, "ISTD": 1.0}
        return workbook, worksheet, provider

    def test_abundance_with_internal_standard(self, setup_base_mocks):
        """Tests 'nmol' and 'Relative' headers and math."""
        workbook, worksheet, provider = setup_base_mocks
        exporter = MagicMock()
        exporter.internal_standard_compound = "ISTD"
        exporter.internal_standard_reference_isotope = 0

        write(workbook, exporter, None, 0, 100, provider=provider)

        # Skip col 0 ('Units') and col 1 (None)
        unit_calls = [
            c.args[2]
            for c in worksheet.write.call_args_list
            if c.args[0] == 4 and c.args[1] >= 2
        ]
        assert unit_calls == ["nmol", "nmol", "Relative"]

        # Math check (Row 5)
        val_calls = [
            c.args[2]
            for c in worksheet.write.call_args_list
            if c.args[0] == 5 and c.args[1] >= 2
        ]
        assert val_calls[1] == pytest.approx(2.5)  # Absolute
        assert val_calls[2] == pytest.approx(5.0)  # Relative

    def test_abundance_with_labelled_internal_standard_reference_peak(self, setup_base_mocks):
        """Tests internal standard normalization using M+N reference peak."""
        workbook, worksheet, provider = setup_base_mocks
        exporter = MagicMock()
        exporter.internal_standard_compound = "ISTD"
        exporter.internal_standard_reference_isotope = 1

        # Override sample data so IS has M0 and M1.
        # Total signal for compound is sum of all isotopologues.
        provider.get_sample_corrected_data.return_value = {
            "ISTD": [1000.0, 100.0],
            "Cmp_Absolute": [500.0],
            "Cmp_Relative": [500.0],
        }

        write(workbook, exporter, None, 0, 100, provider=provider)

        # Math check (Row 5)
        val_calls = [
            c.args[2]
            for c in worksheet.write.call_args_list
            if c.args[0] == 5 and c.args[1] >= 2
        ]

        # Using M+1 internal standard signal (100.0):
        # Absolute: total=500, Amount_IS=10, IS_ref=100, MRRF=2 => 500*(10/100)*(1/2)=25
        assert val_calls[1] == pytest.approx(25.0)
        # Relative: total=500, Amount_IS=10, IS_ref=100 => 500/100*10=50
        assert val_calls[2] == pytest.approx(50.0)

    def test_abundance_no_internal_standard(self, setup_base_mocks):
        """Tests 'Peak Area' mode when no standard is selected."""
        workbook, worksheet, provider = setup_base_mocks
        exporter = MagicMock()
        exporter.internal_standard_compound = None  # Trigger no-standard mode

        # Use the names that match the signal data in setup_base_mocks
        compounds = [
            {
                "compound_name": "Cmp_Absolute",
                "amount_in_std_mix": 1.0,
                "mass0": 100,
                "retention_time": 5.0,
            },
            {
                "compound_name": "Cmp_Relative",
                "amount_in_std_mix": 0.0,
                "mass0": 110,
                "retention_time": 5.5,
            },
        ]
        provider.get_all_compounds.return_value = compounds

        write(workbook, exporter, None, 0, 100, provider=provider)

        # Assert Headers (Row 4): Should all be "Peak Area"
        unit_calls = [
            c.args[2]
            for c in worksheet.write.call_args_list
            if c.args[0] == 4 and c.args[1] >= 2
        ]
        assert all(u == "Peak Area" for u in unit_calls)

        # Assert Values (Row 5): Signals are 500.0 in the fixture
        val_calls = [
            c.args[2]
            for c in worksheet.write.call_args_list
            if c.args[0] == 5 and c.args[1] >= 2
        ]
        assert val_calls[0] == 500.0  # Sum of corrected areas for Cmp_Absolute
        assert val_calls[1] == 500.0  # Sum of corrected areas for Cmp_Relative

    def test_missing_internal_standard_signal_leaves_cells_empty(self, setup_base_mocks):
        """A sample with no IS signal cannot be calibrated, so it must not read as 0."""
        workbook, worksheet, provider = setup_base_mocks
        exporter = MagicMock()
        exporter.internal_standard_compound = "ISTD"
        exporter.internal_standard_reference_isotope = 0
        provider.get_sample_corrected_data.return_value = {
            "ISTD": [0.0],
            "Cmp_Absolute": [500.0],
            "Cmp_Relative": [500.0],
        }

        write(workbook, exporter, None, 0, 100, provider=provider)

        val_calls = [
            c.args[2]
            for c in worksheet.write.call_args_list
            if c.args[0] == 5 and c.args[1] >= 2
        ]
        assert val_calls == [10.0, None, None]

    def test_missing_internal_standard_amount_leaves_cells_empty_outside_mm(
        self, setup_base_mocks
    ):
        """No int_std_amount: MM samples still calibrate from amount_in_std_mix, others are empty."""
        workbook, worksheet, provider = setup_base_mocks
        exporter = MagicMock()
        exporter.internal_standard_compound = "ISTD"
        exporter.internal_standard_reference_isotope = 0
        provider.get_all_samples.return_value = ["MM_01", "S1"]
        provider.resolve_mm_samples.return_value = ["MM_01"]
        provider.get_all_compounds.return_value[0].update(
            int_std_amount=None, mm_files="MM"
        )

        write(workbook, exporter, None, 0, 100, provider=provider)

        rows = {
            row: [
                c.args[2]
                for c in worksheet.write.call_args_list
                if c.args[0] == row and c.args[1] >= 2
            ]
            for row in (5, 6)
        }
        assert rows == {5: [1.0, 0.25, 0.5], 6: [None, None, None]}


def _unlabelled_abundance_rows(bulk, int_std_amount=5.0, mm_samples=()):
    from io import BytesIO
    from types import SimpleNamespace

    import xlsxwriter
    from openpyxl import load_workbook

    from manic.sheet_generators import unlabelled_results

    provider = MagicMock()
    provider.get_mrrf_values.return_value = {"ISTD": 1.0, "Met": 2.0}
    provider.resolve_mm_samples.return_value = list(mm_samples)
    exporter = SimpleNamespace(_provider=provider, internal_standard_compound="ISTD")
    compounds = [
        SimpleNamespace(
            compound_name=name,
            baseline_correction=True,
            retention_time=1.0,
            analysis_channels=[SimpleNamespace(mz=100.0)],
        )
        for name in ("ISTD", "Met")
    ]
    meta = {
        name: {"amount_in_std_mix": 1.0, "int_std_amount": int_std_amount, "mm_files": ""}
        for name in ("ISTD", "Met")
    }
    out = BytesIO()
    workbook = xlsxwriter.Workbook(out, {"in_memory": True})
    unlabelled_results._write_abundances(
        workbook, exporter, list(bulk), compounds, list(meta.values()), meta, bulk, None
    )
    workbook.close()
    out.seek(0)
    sheet = load_workbook(out)["Abundances"]
    return list(sheet.iter_rows(min_row=5, values_only=True))


def test_unlabelled_missing_internal_standard_signal_leaves_cells_empty():
    rows = _unlabelled_abundance_rows(
        {
            "S_ok": {"ISTD": [1000.0], "Met": [4000.0]},
            "S_noIS": {"ISTD": [0.0], "Met": [4000.0]},
        }
    )
    assert rows == [(None, "S_ok", 5, 10), (None, "S_noIS", 5, None)]


def test_unlabelled_mm_sample_uses_amount_in_std_mix_without_int_std_amount():
    rows = _unlabelled_abundance_rows(
        {
            "MM_01": {"ISTD": [1000.0], "Met": [4000.0]},
            "S1": {"ISTD": [1000.0], "Met": [4000.0]},
        },
        int_std_amount=None,
        mm_samples=["MM_01"],
    )
    assert rows == [(None, "MM_01", 1, 2), (None, "S1", None, None)]
